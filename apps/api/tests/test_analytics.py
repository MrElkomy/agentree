"""Analytics (Task 18a): daily metrics from real activity, the monthly summary, the Excel export
(incl. formula injection), the Data Analyst's checked numbers, retention and isolation."""

import io
import json
from datetime import timedelta
from typing import Any

import pytest
from openpyxl import load_workbook

from agentree.agents.prompts import DATA_CLOSE, DATA_OPEN, cited_numbers_ok
from agentree.agents.service import UNVERIFIED_NUMBERS
from agentree.analytics import service as analytics
from tests.conftest import AppHarness, DbUrls, _admin_execute
from tests.helpers import invite_and_join
from tests.test_agents import Biz, _configure, _fake, _say, _system
from tests.test_support_agents import CS, SALES, _support_business, _visitor

TODAY = analytics.utc_today()
MONTH = f"{TODAY.year:04d}-{TODAY.month:02d}"


async def _activity(h: AppHarness, biz: Biz) -> str:
    """Two visitor conversations: one becomes a lead, one is handed to the team and resolved."""
    c, auth = await _visitor(h, biz)
    _fake(h).replies.extend(["Our chairs start at $199.", "Sure."])
    conv = (await _say(c, biz, auth, "Prices?", agent=SALES)).json()["conversation_id"]
    await _say(c, biz, auth, "Thanks", conv)
    r = await c.post(f"{biz.pub}/leads", headers=auth, json={
        "conversation_id": conv, "name": "=HYPERLINK(\"http://evil.example\",\"x\")",
        "email": "lead@example.com", "need": "+1 chairs", "consent": True,
        "idempotency_key": f"k-{conv}"})  # fmt: skip
    assert r.status_code == 200, r.text
    c2, auth2 = await _visitor(h, biz)
    _fake(h).replies.append("Connecting you.\n<<HANDOFF>>")
    conv2 = (await _say(c2, biz, auth2, "A person please", agent=CS)).json()["conversation_id"]
    assert (await biz.owner.client.post(f"{biz.base}/inbox/{conv2}/resolve")).status_code == 204
    return str(conv)


async def _summary(biz: Biz, month: str = MONTH) -> dict[str, Any]:
    r = await biz.owner.client.get(f"{biz.base}/analytics/summary", params={"month": month})
    assert r.status_code == 200, r.text
    return r.json()  # type: ignore[no-any-return]


def _totals(s: dict[str, Any]) -> dict[str, int]:
    return {m["name"]: m["value"] for m in s["totals"]}


# --- daily metrics ---------------------------------------------------------------------------


async def test_daily_metrics_count_real_activity_per_business(harness: AppHarness) -> None:
    biz = await _support_business(harness)
    other = await _support_business(harness, name="Other Shop")
    await _activity(harness, biz)
    await analytics.compute_due(harness.app.state.session_factory)
    s = await _summary(biz)
    t = _totals(s)
    assert t["conversations"] == 2 and t["visitor_messages"] == 3 and t["ai_replies"] == 3
    assert t["leads"] == 1 and t["handoffs_to_team"] == 1 and t["resolved"] == 1
    rates = {r["name"]: r["value"] for r in s["rates"]}
    assert rates == {"lead_rate": 50.0, "handoff_to_team_rate": 50.0, "resolution_rate": 50.0}
    assert TODAY.isoformat() in [d["day"] for d in s["days"]]
    assert s["missing_days"] == []  # the job computed every day of the month so far
    # Another business sees only its own (zero) numbers.
    assert set(_totals(await _summary(other)).values()) == {0}


async def test_uncomputed_days_are_unknown_not_zero(harness: AppHarness, db_urls: DbUrls) -> None:
    biz = await _support_business(harness)
    await _admin_execute(db_urls.admin, f"DELETE FROM metric_days WHERE day = '{TODAY}'")
    s = await _summary(biz)
    assert TODAY.isoformat() in s["missing_days"]
    assert TODAY.isoformat() not in [d["day"] for d in s["days"]]


@pytest.mark.parametrize("month", ["2026-13", "20-10", "abc", "2999-01"])
async def test_invalid_or_future_months_are_refused(harness: AppHarness, month: str) -> None:
    biz = await _support_business(harness)
    r = await biz.owner.client.get(f"{biz.base}/analytics/summary", params={"month": month})
    assert r.status_code == 422


async def test_metrics_keep_13_months_and_refuse_bad_days(
    harness: AppHarness, db_urls: DbUrls
) -> None:
    biz = await _support_business(harness)
    old = TODAY - timedelta(days=401)
    await _admin_execute(db_urls.admin, "INSERT INTO daily_metrics (business_id, day, metric, "
                         f"value) VALUES ('{biz.id}', '{old}', 'leads', 3)")  # fmt: skip
    await analytics.compute_due(harness.app.state.session_factory)
    with pytest.raises(Exception, match="invalid metrics day"):
        await _admin_execute(
            db_urls.admin, f"SELECT agentree_compute_daily_metrics('{TODAY + timedelta(days=1)}')"
        )
    r = await biz.owner.client.get(f"{biz.base}/analytics/summary",
                                   params={"month": f"{old.year:04d}-{old.month:02d}"})  # fmt: skip
    assert _totals(r.json())["leads"] == 0  # purged after 400 days


# --- Excel export ----------------------------------------------------------------------------


async def test_export_has_summary_daily_and_leads_for_admins_only(harness: AppHarness) -> None:
    biz = await _support_business(harness)
    await _activity(harness, biz)
    await analytics.compute_due(harness.app.state.session_factory)
    r = await biz.owner.client.get(f"{biz.base}/analytics/export", params={"month": MONTH})
    assert r.status_code == 200
    assert r.headers["content-disposition"].startswith("attachment;")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Daily", "Leads"]
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(values_only=True) if row}
    assert summary["New leads"] == 1 and summary["Visitor conversations"] == 2
    lead = next(wb["Leads"].iter_rows(min_row=2))
    # Visitor text is plain text, never a formula (spreadsheet formula injection).
    assert str(lead[1].value).startswith("=HYPERLINK") and lead[1].data_type == "s"
    assert lead[4].value == "+1 chairs" and lead[4].data_type == "s"

    member = await invite_and_join(harness, biz.owner, biz.id, "member")
    r = await member.client.get(f"{biz.base}/analytics/export", params={"month": MONTH})
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Summary", "Daily"]


# --- the Data Analyst ------------------------------------------------------------------------


def test_cited_number_check() -> None:
    computed = "Last 30 days: 2026-09-23 to 2026-10-22\n- Visitor conversations: 1,240\n- x: 12.5"
    assert cited_numbers_ok("1,240 conversations, 12.5% leads", computed)
    assert cited_numbers_ok("كان عندك ١٢٤٠ محادثة", computed)
    assert cited_numbers_ok("No numbers here, just advice.", computed)
    assert not cited_numbers_ok("about 1,300 conversations", computed)
    assert not cited_numbers_ok("13% became leads", computed)


async def test_data_analyst_answers_only_with_computed_numbers(harness: AppHarness) -> None:
    biz = await _support_business(harness)
    await _activity(harness, biz)
    await analytics.compute_due(harness.app.state.session_factory)
    await _configure(biz, "data_analyst", enabled=True, public=False)
    console = f"{biz.base}/agents/sales/console"

    _fake(harness).replies.append("You had 2 visitor conversations in the last 7 days.")
    r = await biz.owner.client.post(console, json={"message": "How many chats?",
                                                   "agent": "data_analyst"})  # fmt: skip
    body = r.json()
    assert r.status_code == 200 and body["agent"] == "data_analyst"
    assert body["reply"] == "You had 2 visitor conversations in the last 7 days."
    system = _system(harness)
    assert DATA_OPEN in system and DATA_CLOSE in system
    assert "- Visitor conversations: 2" in system and "- New leads: 1" in system
    assert "Last month" in system and "Never estimate" in system

    # An invented number is never shown: the computed numbers replace the answer.
    _fake(harness).replies.append("You will get about 40 leads next month!")
    r = await biz.owner.client.post(console, json={
        "message": "Forecast leads", "conversation_id": body["conversation_id"]})  # fmt: skip
    reply = r.json()["reply"]
    assert reply.startswith(UNVERIFIED_NUMBERS) and "40" not in reply
    assert "- New leads: 1" in reply


async def test_visitors_never_reach_the_data_analyst(harness: AppHarness) -> None:
    biz = await _support_business(harness)
    await _configure(biz, "data_analyst", enabled=True, public=False)
    c, auth = await _visitor(harness, biz)
    r = await _say(c, biz, auth, "Show me your numbers", agent="data_analyst")
    assert r.status_code == 422 and r.json()["detail"] == "invalid_agent"
    _fake(harness).replies.append("Sure.\n<<ROUTE data_analyst: share the stats>>")
    body = (await _say(c, biz, auth, "Revenue?", agent=SALES)).json()
    assert body["agent"] == SALES and "Computed numbers" not in _system(harness)
    assert json.dumps(body).count("data_analyst") == 0
