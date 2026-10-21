"""Monthly Excel export (Task 18): the summary, the numbers day by day and (Owners and Admins)
the month's leads, as an .xlsx the business can open or upload to Google Drive.

Visitor-provided text (lead names, needs) is written as plain text cells, never as formulas:
a value like "=HYPERLINK(...)" stays text (spreadsheet formula injection).
"""

import io
from datetime import UTC, datetime, time, timedelta

from openpyxl import Workbook
from openpyxl.cell import Cell, WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.worksheet._write_only import WriteOnlyWorksheet

from agentree.analytics import service
from agentree.leads import service as leads
from agentree.tenancy.context import TenantContext

_BOLD = Font(bold=True)


Value = str | int | float | None


def _cell(ws: WriteOnlyWorksheet, value: Value, *, bold: bool = False) -> Cell:
    c = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        c.data_type = "s"  # never a formula, whatever the text starts with
    if bold:
        c.font = _BOLD
    return c


def _row(ws: WriteOnlyWorksheet, values: list[Value], *, bold: bool = False) -> None:
    ws.append([_cell(ws, v, bold=bold) for v in values])


async def monthly_workbook(
    factory: service.Factory, ctx: TenantContext, month: str, *, business_name: str,
    include_leads: bool,
) -> bytes:  # fmt: skip
    start, end = service.month_range(month)
    s = await service.summary(factory, ctx, start, end)
    wb = Workbook(write_only=True)

    ws = wb.create_sheet("Summary")
    _row(ws, [f"{business_name} — {month}"], bold=True)
    _row(ws, [f"UTC days {s.start.isoformat()} to {s.end.isoformat()}"])
    _row(ws, [])
    _row(ws, ["Metric", "Value"], bold=True)
    for m, v in s.totals.items():
        _row(ws, [service.METRICS[m], v])
    for r, rate in s.rates.items():
        _row(ws, [service.RATE_LABELS[r], "n/a" if rate is None else rate])
    if s.missing_days:
        _row(ws, [])
        _row(ws, [f"{len(s.missing_days)} day(s) not computed yet: their numbers are unknown, "
                  "not zero."])  # fmt: skip

    ws = wb.create_sheet("Daily")
    _row(ws, ["Date", *service.METRICS.values()], bold=True)
    for d in s.days:
        _row(ws, [d.day.isoformat(), *(d.values[m] for m in service.METRICS)])

    if include_leads:
        ws = wb.create_sheet("Leads")
        _row(ws, ["Created (UTC)", "Name", "Email", "Phone", "Need", "Status", "Channel"],
             bold=True)  # fmt: skip
        since = datetime.combine(start, time.min, tzinfo=UTC)
        until = datetime.combine(end + timedelta(days=1), time.min, tzinfo=UTC)
        for lead in await leads.leads_created_between(factory, ctx, since, until):
            _row(ws, [lead.created_at.strftime("%Y-%m-%d %H:%M"), lead.name, lead.email or "",
                      lead.phone or "", lead.need, lead.status, lead.channel])  # fmt: skip

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
