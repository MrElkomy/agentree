"""Uploaded CSV/XLSX as a queryable table (Task 18b). Runs only in a limited child process.

    python -m agentree.knowledge.tables      # stdin: one JSON header line, then the file bytes

Header `{"mode": "profile", "kind": "csv"}` returns the columns (names, inferred types) and the
row count, never cell values. Header `{"mode": "query", "kind": "xlsx", "sql": "SELECT ..."}`
loads the file into an in-memory SQLite table named `data` and runs one read-only statement.

The sandbox, besides the child's memory/CPU/file-write limits and missing secrets
(`core.sandbox`): a single SELECT/WITH statement; an SQLite authorizer that allows reading only
the `data` table and calling only allow-listed functions (no ATTACH, no PRAGMA, no writes, no
schema reads, no recursion); a progress handler that stops long queries; result rows capped.
Python's sqlite3 has extension loading off, and the database is in memory only. Workbook XML is
parsed by openpyxl, which uses defusedxml (installed) against entity attacks.

SQLite rather than DuckDB (the roadmap's first idea): it ships with Python, needs far less
memory on the shared host and keeps the same guarantees (read-only, no files, no network).
"""

import csv
import io
import json
import re
import sqlite3
import sys
from datetime import date, datetime, time
from typing import Any

MAX_ROWS = 200_000
MAX_COLUMNS = 100
MAX_CELL = 1000
MAX_RESULT_ROWS = 200
MAX_SQL = 4000
MAX_STEPS = 20_000_000  # SQLite VM steps (progress handler every 10,000)

ALLOWED_FUNCTIONS = frozenset({
    "count", "sum", "total", "avg", "min", "max", "round", "abs", "length", "lower", "upper",
    "trim", "ltrim", "rtrim", "substr", "substring", "replace", "instr", "coalesce", "ifnull",
    "nullif", "iif", "date", "time", "datetime", "strftime", "julianday", "cast", "typeof",
    "group_concat", "like", "glob", "printf", "format",
})  # fmt: skip

_NAME = re.compile(r"[^\w]+", re.UNICODE)


class TableError(ValueError):
    """`str(exc)` is a stable code (no_rows, too_many_columns, invalid_sql, query_failed...)."""


def _column_names(header: list[Any]) -> list[str]:
    names: list[str] = []
    for i, raw in enumerate(header, start=1):
        name = _NAME.sub("_", str(raw or "").strip().lower()).strip("_")[:40] or f"column_{i}"
        if name[0].isdigit():
            name = f"c_{name}"[:40]
        base, n = name, 2
        while name in names:
            name, n = f"{base}_{n}", n + 1
        names.append(name)
    return names


def _cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int | float):
        return value
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    text = str(value).strip()[:MAX_CELL]
    if not text:
        return None
    number = text.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+(\.\d+)?", text) else text
    if re.fullmatch(r"-?\d+", number):
        return int(number)
    if re.fullmatch(r"-?\d*\.\d+", number):
        return float(number)
    return text


def _rows_csv(data: bytes) -> list[list[Any]]:
    text = data.decode("utf-8-sig")
    try:
        dialect: Any = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows: list[list[Any]] = []
    for row in csv.reader(io.StringIO(text), dialect):
        rows.append(list(row))
        if len(rows) > MAX_ROWS + 1:
            raise TableError("too_many_rows")
    return rows


def _rows_xlsx(data: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows: list[list[Any]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
        if len(rows) > MAX_ROWS + 1:
            raise TableError("too_many_rows")
    return rows


def load(kind: str, data: bytes) -> tuple[list[str], list[str], list[list[Any]]]:
    """(column names, column types, rows) of the file's first sheet; header = first row."""
    raw = _rows_csv(data) if kind == "csv" else _rows_xlsx(data)
    raw = [r for r in raw if any(c not in (None, "") for c in r)]
    if len(raw) < 2:
        raise TableError("no_rows")
    width = max(len(r) for r in raw)
    if width > MAX_COLUMNS:
        raise TableError("too_many_columns")
    names = _column_names(raw[0] + [None] * (width - len(raw[0])))
    rows = [[_cell(v) for v in r] + [None] * (width - len(r)) for r in raw[1:]]
    types = []
    for i in range(width):
        values = [r[i] for r in rows if r[i] is not None]
        if values and all(isinstance(v, int) for v in values):
            types.append("integer")
        elif values and all(isinstance(v, int | float) for v in values):
            types.append("real")
        else:
            types.append("text")
            for r in rows:
                if r[i] is not None and not isinstance(r[i], str):
                    r[i] = str(r[i])
    return names, types, rows


def _authorizer(action: int, arg1: str | None, arg2: str | None, _db: str | None,
                _trigger: str | None) -> int:  # fmt: skip
    if action == sqlite3.SQLITE_SELECT:
        return sqlite3.SQLITE_OK
    if action == sqlite3.SQLITE_READ:
        return sqlite3.SQLITE_OK if arg1 == "data" else sqlite3.SQLITE_DENY
    if action == sqlite3.SQLITE_FUNCTION:
        return (
            sqlite3.SQLITE_OK if (arg2 or "").lower() in ALLOWED_FUNCTIONS else sqlite3.SQLITE_DENY
        )
    return sqlite3.SQLITE_DENY  # writes, ATTACH, PRAGMA, schema, recursion, transactions...


def validate_sql(sql: str) -> str:
    statement = sql.strip().rstrip(";").strip()
    if not statement or len(statement) > MAX_SQL or ";" in statement:
        raise TableError("invalid_sql")
    if not re.match(r"(?is)^(select|with)\b", statement):
        raise TableError("invalid_sql")
    return statement


def query(kind: str, data: bytes, sql: str) -> dict[str, Any]:
    statement = validate_sql(sql)
    names, types, rows = load(kind, data)
    conn = sqlite3.connect(":memory:")
    cols = ", ".join(f'"{n}" {t.upper()}' for n, t in zip(names, types, strict=True))
    # Column names are reduced to word characters by _column_names, then quoted; values are
    # bound parameters. Nothing from the file is ever spliced into SQL as text.
    conn.execute(f"CREATE TABLE data ({cols})")
    marks = ", ".join("?" * len(names))
    conn.executemany(f"INSERT INTO data VALUES ({marks})", rows)  # noqa: S608
    conn.setlimit(sqlite3.SQLITE_LIMIT_ATTACHED, 0)
    conn.setlimit(sqlite3.SQLITE_LIMIT_SQL_LENGTH, MAX_SQL + 100)
    conn.set_authorizer(_authorizer)
    steps = [0]

    def progress() -> int:
        steps[0] += 10_000
        return 1 if steps[0] > MAX_STEPS else 0  # non-zero aborts the query

    conn.set_progress_handler(progress, 10_000)
    try:
        cur = conn.execute(statement)
        result = cur.fetchmany(MAX_RESULT_ROWS + 1)
    except sqlite3.DatabaseError as exc:
        code = "query_too_long" if steps[0] > MAX_STEPS else "query_failed"
        raise TableError(code) from exc
    columns = [d[0] for d in cur.description or []]
    return {"columns": columns, "rows": [list(r) for r in result[:MAX_RESULT_ROWS]],
            "truncated": len(result) > MAX_RESULT_ROWS}  # fmt: skip


def profile(kind: str, data: bytes) -> dict[str, Any]:
    names, types, rows = load(kind, data)
    return {"columns": [{"name": n, "type": t} for n, t in zip(names, types, strict=True)],
            "rows": len(rows)}  # fmt: skip


def main() -> None:
    stream = sys.stdin.buffer
    header = json.loads(stream.readline(MAX_SQL * 4))
    data = stream.read()
    try:
        if header["mode"] == "profile":
            out: dict[str, Any] = {"ok": True, **profile(header["kind"], data)}
        else:
            out = {"ok": True, **query(header["kind"], data, str(header["sql"]))}
    except TableError as exc:
        out = {"ok": False, "error": str(exc)}
    except Exception:  # unreadable file or anything unexpected: no details leave the child
        out = {"ok": False, "error": "damaged_file"}
    sys.stdout.write(json.dumps(out, default=str))


if __name__ == "__main__":
    main()
